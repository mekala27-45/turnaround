"""The Excel workbook: one sheet per chapter's table, the monthly marts, and a summary that computes.

The summary sheet does not hold numbers. It holds a year (a cell the reader changes, with a list of
the years in the data) and formulas over the carrier by month table and the inherited month table, so
a manager can change the year and watch the on time rate, the padding, the cancellation rate, the
reported cause shares and the inherited share move. The ranking cells read the chapter 5 table with
INDEX and MATCH. Every input range is a named range, so the formulas read as words.

The workbook is written with openpyxl, which stores formulas but no cached values; the test opens it
in LibreOffice, recalculates, and compares the summary with the manifest for the default year.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from turnaround_core.manifest import Manifest
from turnaround_core.statements import STATEMENT

HEADER_FILL = PatternFill("solid", fgColor="E6E2DB")
INK = "141210"
CARRIER_MONTH_COLUMNS = (
    "carrier",
    "year",
    "month",
    "scheduled",
    "cancelled",
    "diverted",
    "flown",
    "on_time",
    "arr_delay_sum",
    "padding_sum",
    "padding_flights",
    "cause_late_aircraft_sum",
    "cause_carrier_sum",
    "cause_nas_sum",
    "cause_weather_sum",
    "cause_security_sum",
)
INHERITED_COLUMNS = ("carrier", "year", "month", "inherited_minutes_sum", "arr_delay_pos_sum", "flown")
FORMATS = {
    "pct0": "0%",
    "pct1": "0.0%",
    "pct2": "0.00%",
    "pts1": '+0.0" pts";-0.0" pts"',
    "min0": '0" min"',
    "min1": '0.0" min"',
    "min2": '0.00" min"',
    "smin1": '+0.0" min";-0.0" min"',
    "smin2": '+0.00" min";-0.00" min"',
    "int": "#,##0",
    "float1": "0.0",
    "float2": "0.00",
    "float3": "0.000",
}


@dataclass(frozen=True)
class SummaryCell:
    label: str
    formula: str
    fmt: str
    metric: str | None  # the metric layer name it should equal for the default year, when there is one


def _header(ws: Worksheet, columns: Sequence[str], row: int = 1) -> None:
    for j, name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=j, value=name)
        cell.font = Font(bold=True, color=INK)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left" if j == 1 else "right")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _frame_sheet(wb: Workbook, title: str, frame: pl.DataFrame, columns: Sequence[str]) -> Worksheet:
    ws = wb.create_sheet(title)
    _header(ws, columns)
    for i, row in enumerate(frame.select(columns).iter_rows(), start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)
    for j, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, len(name) + 2)
    return ws


def _name_columns(wb: Workbook, ws: Worksheet, prefix: str, columns: Sequence[str], rows: int) -> None:
    for j, name in enumerate(columns, start=1):
        letter = get_column_letter(j)
        ref = f"'{ws.title}'!${letter}$2:${letter}${rows + 1}"
        wb.defined_names[f"{prefix}_{name}"] = DefinedName(f"{prefix}_{name}", attr_text=ref)


def _table_sheet(wb: Workbook, manifest: Manifest, key: str, title: str) -> Worksheet | None:
    entry = manifest.tables.get(key)
    if entry is None:
        return None
    ws = wb.create_sheet(title[:31])
    _header(ws, entry.columns)
    for i, row in enumerate(entry.rows, start=2):
        for j, (value, fmt) in enumerate(zip(row, entry.formats, strict=True), start=1):
            cell = ws.cell(row=i, column=j, value=value)
            if fmt in FORMATS and isinstance(value, int | float):
                cell.number_format = FORMATS[fmt]
    for j, name in enumerate(entry.columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(48, len(name) + 4))
    ws.cell(row=len(entry.rows) + 3, column=1, value=manifest.label(key))
    return ws


def summary_cells() -> list[SummaryCell]:
    """The summary's formulas. YEAR is the named input cell."""

    def ratio(numerator: str, denominator: str) -> str:
        return f"=SUMIFS(cm_{numerator},cm_year,Year)/SUMIFS(cm_{denominator},cm_year,Year)"

    causes = "(SUMIFS(cm_cause_late_aircraft_sum,cm_year,Year)+SUMIFS(cm_cause_carrier_sum,cm_year,Year)+SUMIFS(cm_cause_nas_sum,cm_year,Year)+SUMIFS(cm_cause_weather_sum,cm_year,Year)+SUMIFS(cm_cause_security_sum,cm_year,Year))"
    return [
        SummaryCell("Scheduled flights", "=SUMIFS(cm_scheduled,cm_year,Year)", "int", None),
        SummaryCell("Flights flown", "=SUMIFS(cm_flown,cm_year,Year)", "int", None),
        SummaryCell(
            "On time rate (arrival less than 15 minutes late)",
            ratio("on_time", "flown"),
            "pct1",
            "on_time_rate",
        ),
        SummaryCell("Cancellation rate", ratio("cancelled", "scheduled"), "pct2", "cancellation_rate"),
        SummaryCell("Mean arrival delay", ratio("arr_delay_sum", "flown"), "min1", "mean_arrival_delay"),
        SummaryCell(
            "Mean padding per flight", ratio("padding_sum", "padding_flights"), "min1", "mean_padding"
        ),
        SummaryCell(
            "Reported late aircraft share of cause minutes",
            f"=SUMIFS(cm_cause_late_aircraft_sum,cm_year,Year)/{causes}",
            "pct1",
            "reported_late_aircraft_share",
        ),
        SummaryCell(
            "Reported weather share of cause minutes",
            f"=SUMIFS(cm_cause_weather_sum,cm_year,Year)/{causes}",
            "pct1",
            "reported_weather_share",
        ),
        SummaryCell(
            "Reported air traffic system share of cause minutes",
            f"=SUMIFS(cm_cause_nas_sum,cm_year,Year)/{causes}",
            "pct1",
            "reported_nas_share",
        ),
        SummaryCell(
            "Inherited share of arrival delay minutes (reporting years only)",
            '=IFERROR(SUMIFS(im_inherited_minutes_sum,im_year,Year)/SUMIFS(im_arr_delay_pos_sum,im_year,Year),"not a reporting year")',
            "pct1",
            "inherited_share",
        ),
        SummaryCell(
            "Least delayed carrier once adjusted (chapter 5)",
            "=INDEX(rank_carrier,MATCH(1,rank_adjusted,0))",
            "text",
            None,
        ),
        SummaryCell(
            "Least delayed carrier on the raw average (chapter 5)",
            "=INDEX(rank_carrier,MATCH(1,rank_raw,0))",
            "text",
            None,
        ),
    ]


def build(
    out: Path,
    manifest: Manifest,
    carrier_month: pl.DataFrame,
    inherited_month: pl.DataFrame,
    *,
    default_year: int,
    chapter_tables: Sequence[tuple[str, str]],
) -> dict[str, Any]:
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    summary.title = "Summary"
    cm = carrier_month.sort(["carrier", "year", "month"])
    im = inherited_month.sort(["carrier", "year", "month"])
    cm_ws = _frame_sheet(wb, "carrier_month", cm, CARRIER_MONTH_COLUMNS)
    _name_columns(wb, cm_ws, "cm", CARRIER_MONTH_COLUMNS, cm.height)
    im_ws = _frame_sheet(wb, "inherited_month", im, INHERITED_COLUMNS)
    _name_columns(wb, im_ws, "im", INHERITED_COLUMNS, max(im.height, 1))

    ranking = manifest.tables["ch5.ranking"]
    rank_ws = wb.create_sheet("ranking")
    rank_cols = ["carrier", "raw_rank", "adjusted_rank", "raw_effect", "adjusted_effect"]
    _header(rank_ws, rank_cols)
    idx = {
        c: ranking.columns.index(c)
        for c in ("Carrier", "Raw rank", "Adjusted rank", "Raw effect", "Adjusted effect")
    }
    for i, row in enumerate(ranking.rows, start=2):
        rank_ws.cell(row=i, column=1, value=row[idx["Carrier"]])
        rank_ws.cell(row=i, column=2, value=row[idx["Raw rank"]])
        rank_ws.cell(row=i, column=3, value=row[idx["Adjusted rank"]])
        rank_ws.cell(row=i, column=4, value=row[idx["Raw effect"]]).number_format = FORMATS["smin2"]
        rank_ws.cell(row=i, column=5, value=row[idx["Adjusted effect"]]).number_format = FORMATS["smin2"]
    n_rank = len(ranking.rows)
    for j, name in enumerate(("rank_carrier", "rank_raw", "rank_adjusted"), start=1):
        letter = get_column_letter(j)
        wb.defined_names[name] = DefinedName(name, attr_text=f"'ranking'!${letter}$2:${letter}${n_rank + 1}")

    years = sorted({int(y) for y in cm["year"].to_list()})
    years_ws = wb.create_sheet("years")
    for i, y in enumerate(years, start=1):
        years_ws.cell(row=i, column=1, value=y)
    years_ws.sheet_state = "hidden"

    summary["A1"] = "turnaround: the summary computes from the tables in this workbook"
    summary["A1"].font = Font(bold=True, size=14, color=INK)
    summary["A2"] = STATEMENT
    summary["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells("A2:C2")
    summary.row_dimensions[2].height = 60
    summary["A4"] = "Year (change me)"
    summary["A4"].font = Font(bold=True)
    summary["B4"] = default_year
    summary["B4"].fill = PatternFill("solid", fgColor="FCE9DF")
    wb.defined_names["Year"] = DefinedName("Year", attr_text="'Summary'!$B$4")
    validation = DataValidation(type="list", formula1=f"='years'!$A$1:$A${len(years)}", allow_blank=False)
    summary.add_data_validation(validation)
    validation.add("B4")
    cells = summary_cells()
    rows: dict[str, int] = {}
    for i, c in enumerate(cells, start=6):
        summary.cell(row=i, column=1, value=c.label)
        cell = summary.cell(row=i, column=2, value=c.formula)
        if c.fmt in FORMATS:
            cell.number_format = FORMATS[c.fmt]
        summary.cell(row=i, column=3, value=c.metric or "")
        rows[c.label] = i
    summary["C5"] = "metric layer name"
    summary["C5"].font = Font(italic=True, color="5A5651")
    summary.column_dimensions["A"].width = 64
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 30

    sheets = 0
    for key, title in chapter_tables:
        if _table_sheet(wb, manifest, key, title) is not None:
            sheets += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {
        "sheets": len(wb.sheetnames),
        "chapter_sheets": sheets,
        "formulas": len(cells),
        "named_ranges": len(wb.defined_names),
        "default_year": default_year,
        "rows": rows,
    }


def recalculate(path: Path, out_dir: Path) -> Path:
    """Open the workbook in LibreOffice headless, recalculate every formula, save a copy with values."""
    import shutil
    import subprocess

    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        raise FileNotFoundError("LibreOffice is not installed")
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [
            office,
            "--headless",
            "--norestore",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            str(out_dir),
            str(path),
        ],
        check=True,
        capture_output=True,
        timeout=180,
    )
    return out_dir / path.name


def read_summary(recalculated: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(recalculated, data_only=True)
    ws = wb["Summary"]
    values: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=6, max_col=3, values_only=True):
        label, value, metric = row
        if label:
            values[str(label)] = {"value": value, "metric": metric or None}
    values["__year__"] = ws["B4"].value
    return values
