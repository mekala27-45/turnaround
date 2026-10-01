"""The workbook's formulas recalculated by LibreOffice against the metric values; the extracts and the
bundle against their sources."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

import numpy as np
import polars as pl
import pytest
from turnaround_core.manifest import Manifest
from turnaround_exports import bi, workbook

from tests._deps import need

ROOT = Path(__file__).resolve().parents[2]


def _carrier_month(seed: int = 3) -> pl.DataFrame:
    rng = np.random.default_rng(seed)
    rows = []
    for carrier in ("AA", "DL", "WN"):
        for year in (2024, 2025):
            for month in range(1, 13):
                scheduled = int(rng.integers(20_000, 40_000))
                cancelled = int(scheduled * rng.uniform(0.005, 0.03))
                diverted = int(scheduled * 0.002)
                flown = scheduled - cancelled - diverted
                causes = rng.uniform(1e5, 5e5, 5)
                rows.append(
                    {
                        "carrier": carrier, "year": year, "month": month, "scheduled": scheduled,
                        "cancelled": cancelled, "diverted": diverted, "flown": flown,
                        "on_time": int(flown * rng.uniform(0.7, 0.85)), "arr_delay_sum": float(flown * rng.uniform(2, 12)),
                        "padding_sum": float(flown * rng.uniform(10, 20)), "padding_flights": flown - 100,
                        "cause_late_aircraft_sum": causes[0], "cause_carrier_sum": causes[1], "cause_nas_sum": causes[2],
                        "cause_weather_sum": causes[3], "cause_security_sum": causes[4],
                    }
                )  # fmt: skip
    return pl.DataFrame(rows)


def _inherited(cm: pl.DataFrame) -> pl.DataFrame:
    return cm.filter(pl.col("year") == 2025).select(
        "carrier",
        "year",
        "month",
        (pl.col("arr_delay_sum") * 0.4).alias("inherited_minutes_sum"),
        (pl.col("arr_delay_sum") * 1.3).alias("arr_delay_pos_sum"),
        "flown",
    )


def _manifest() -> Manifest:
    m = Manifest(as_of="2026-09-30", seed=1)
    m.put_table(
        "ch5.ranking",
        ["Carrier", "Flights", "Raw rank", "Raw effect", "Adjusted rank", "Adjusted effect", "Interval", "Moved"],
        ["text", "int", "int", "smin2", "int", "smin2", "text", "text"],
        [["Delta Air Lines", 10, 2, -1.0, 1, -2.0, "", "up 1"], ["Southwest Airlines", 10, 1, -1.5, 2, -0.5, "", "down 1"]],
        source="real:bts", population="test", origin="test",
    )  # fmt: skip
    m.put_table(
        "ch1.by_year",
        ["Year", "On time"],
        ["text", "pct1"],
        [["2024", 0.8], ["2025", 0.81]],
        source="real:bts",
        population="test",
        origin="test",
    )
    return m


def expected(cm: pl.DataFrame, im: pl.DataFrame, year: int) -> dict[str, float]:
    y = cm.filter(pl.col("year") == year)
    causes = sum(
        float(y[f"cause_{c}_sum"].sum()) for c in ("late_aircraft", "carrier", "nas", "weather", "security")
    )
    yi = im.filter(pl.col("year") == year)
    return {
        "on_time_rate": float(y["on_time"].sum()) / float(y["flown"].sum()),
        "cancellation_rate": float(y["cancelled"].sum()) / float(y["scheduled"].sum()),
        "mean_arrival_delay": float(y["arr_delay_sum"].sum()) / float(y["flown"].sum()),
        "mean_padding": float(y["padding_sum"].sum()) / float(y["padding_flights"].sum()),
        "reported_late_aircraft_share": float(y["cause_late_aircraft_sum"].sum()) / causes,
        "reported_weather_share": float(y["cause_weather_sum"].sum()) / causes,
        "reported_nas_share": float(y["cause_nas_sum"].sum()) / causes,
        "inherited_share": float(yi["inherited_minutes_sum"].sum()) / float(yi["arr_delay_pos_sum"].sum()),
    }


def test_workbook_has_live_formulas_and_named_ranges(tmp_path: Path) -> None:
    cm = _carrier_month()
    built = workbook.build(
        tmp_path / "w.xlsx",
        _manifest(),
        cm,
        _inherited(cm),
        default_year=2025,
        chapter_tables=[("ch1.by_year", "Ch1")],
    )
    from openpyxl import load_workbook

    wb = load_workbook(tmp_path / "w.xlsx")
    formulas = [
        c.value
        for row in wb["Summary"].iter_rows(min_row=6, max_col=2)
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert len(formulas) == built["formulas"] == len(workbook.summary_cells())
    assert "Year" in wb.defined_names and "cm_on_time" in wb.defined_names
    assert "Ch1" in wb.sheetnames


@pytest.mark.libreoffice
def test_recalculated_summary_equals_the_metric_values(tmp_path: Path) -> None:
    need(
        "libreoffice",
        shutil.which("soffice") is not None or shutil.which("libreoffice") is not None,
        "LibreOffice is not installed, so the workbook cannot be recalculated",
    )
    cm = _carrier_month()
    im = _inherited(cm)
    workbook.build(tmp_path / "w.xlsx", _manifest(), cm, im, default_year=2025, chapter_tables=[])
    values = workbook.read_summary(workbook.recalculate(tmp_path / "w.xlsx", tmp_path / "recalc"))
    want = expected(cm, im, 2025)
    checked = 0
    for entry in values.values():
        if isinstance(entry, dict) and entry["metric"] in want:
            assert entry["value"] == pytest.approx(want[entry["metric"]], rel=1e-9)
            checked += 1
    assert checked == len(want)
    assert values["Least delayed carrier once adjusted (chapter 5)"]["value"] == "Delta Air Lines"
    assert values["Least delayed carrier on the raw average (chapter 5)"]["value"] == "Southwest Airlines"


def test_extracts_carry_exactly_the_marts_rows(tmp_path: Path) -> None:
    marts = tmp_path / "marts"
    marts.mkdir()
    cm = _carrier_month()
    cm.write_parquet(marts / "mart_carrier_month.parquet")
    rm = cm.with_columns(pl.lit("ATL-DEN").alias("route"))
    rm.write_parquet(marts / "mart_route_month.parquet")
    counts = bi.tableau(marts, tmp_path / "tableau")
    assert counts == {"mart_route_month": rm.height, "mart_carrier_month": cm.height}
    assert pl.read_csv(tmp_path / "tableau" / "mart_carrier_month.csv").height == cm.height
    years = sorted((tmp_path / "tableau" / "mart_route_month").glob("mart_route_month_*.csv"))
    assert len(years) == rm["year"].n_unique()
    assert sum(pl.read_csv(f).height for f in years) == rm.height
    assert (tmp_path / "tableau" / "workbook_spec.md").read_text().count("|") > 20


def test_bundle_dictionary_lists_every_column(tmp_path: Path) -> None:
    m = _manifest()
    counts = bi.csv_bundle(
        m,
        ["ch5.ranking", "ch1.by_year", "absent.table"],
        tmp_path,
        {"extra": pl.DataFrame({"a": [1], "b": ["x"]})},
    )
    assert counts == {"ch5_ranking": 2, "ch1_by_year": 2, "extra": 1}
    with (tmp_path / "dictionary.csv").open() as handle:
        listed = {(r["file"], r["column"]) for r in csv.DictReader(handle)}
    for name in ("ch5_ranking", "ch1_by_year", "extra"):
        with (tmp_path / f"{name}.csv").open() as handle:
            header = next(csv.reader(handle))
        assert all((f"{name}.csv", column) in listed for column in header)


def test_powerbi_measures_come_from_the_metric_layer(tmp_path: Path) -> None:
    from turnaround_pipeline.metric_layer import load

    metrics = load(ROOT / "metrics" / "metrics.yml")
    n = bi.powerbi(metrics, tmp_path, 2)
    assert n == len(metrics) + 1
    text = (tmp_path / "model.md").read_text()
    assert "DIVIDE(SUM('carrier_month'[on_time])" in text
    assert "inherited_month" in text
